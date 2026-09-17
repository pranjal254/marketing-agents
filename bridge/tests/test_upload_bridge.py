"""Uploaded-brief entry point over HTTP: multipart .docx/.xlsx → parse → the same
intake pipeline as typed text (extraction with provenance, hold-for-review)."""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from shiftai_shared.config import SharedSettings

from c2c_bridge.app import create_app

EXTRACT_OK = json.dumps(
    {
        "fields": {
            "objective": {"value": "Qualified pipeline for the D365 Copilot offering",
                          "quote": "Generate qualified pipeline", "basis": "stated"},
            "business_unit": {"value": "Technology", "quote": "Practice: Dynamics 365",
                              "basis": "inferred"},
            "vertical": {"value": "manufacturing", "quote": "Manufacturing", "basis": "stated"},
            "target_segment": {"value": "type_3", "quote": "$500M+", "basis": "inferred"},
            "offer_topic": {"value": "D365 Copilot marketing collateral",
                            "quote": "Copilot for Marketing Collateral", "basis": "stated"},
            "channels": {"value": ["email", "events"], "quote": "Email (Pardot)",
                         "basis": "stated"},
            "timeline_start": {"value": "2026-10-01", "quote": "Weeks 1-2", "basis": "inferred"},
            "timeline_end": {"value": "2026-11-30", "quote": "Week 9", "basis": "inferred"},
            "owner": {"value": "rs@levelshift.com", "quote": "Owner: RS", "basis": "stated"},
            "budget_flag": {"value": True, "quote": "budget TBD", "basis": "inferred"},
        },
        "notes": "",
    }
)

CLASSIFY_OK = json.dumps(
    {
        "action_class": "route_for_approval",
        "confidence": 0.9,
        "rationale": "clear BU and vertical",
        "classification": {
            "campaign_type": "demand_gen",
            "priority": "high",
            "channel_mix": ["email", "events"],
            "segment_relevance": "type_3",
            "field_rationale": {"business_unit": "practice line"},
        },
        "normalized_fields": {},
    }
)


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    from shiftai_shared.llm import MockLLMProvider

    import c2c_bridge.app as app_mod

    def fake_build_provider(_: SharedSettings) -> MockLLMProvider:
        return MockLLMProvider(
            default=CLASSIFY_OK,
            script=[(lambda u: "Extract campaign brief fields" in u, EXTRACT_OK)],
        )

    monkeypatch.setattr(app_mod, "build_provider", fake_build_provider)
    app = create_app(
        workdir=tmp_path / "run",
        settings=SharedSettings(_env_file=None, LLM_PROVIDER="mock"),
    )
    return TestClient(app)


def _docx_upload() -> tuple[str, io.BytesIO, str]:
    from docx import Document

    document = Document()
    document.add_heading("Campaign Brief: D365 Copilot for Marketing Collateral", level=1)
    document.add_paragraph(
        "Generate qualified pipeline for LevelShift's D365 Copilot offering among "
        "manufacturing accounts of $500M+ revenue. Email (Pardot) and events."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return ("brief.docx", buffer, mime)


def _post_upload(client: TestClient, name: str, buffer: io.BytesIO, mime: str) -> Any:
    return client.post(
        "/api/requests/upload",
        files={"file": (name, buffer, mime)},
        data={"requester": "rs@levelshift.com"},
    )


def test_docx_upload_runs_intake_and_holds_for_review(client: TestClient) -> None:
    response = _post_upload(client, *_docx_upload())
    assert response.status_code == 200, response.text
    outcome = response.json()
    # hold_for_verification defaults true: the drafted brief waits for the requester
    assert outcome["status"] == "draft_review"
    assert outcome["upload"]["kind"] == "docx"
    assert outcome["upload"]["truncated"] is False
    detail = client.get(f"/api/cases/{outcome['case_id']}").json()
    request = detail["summary"]["request"]
    # the document's own words became the intake description, in document order
    assert request["free_text_context"].startswith("# Campaign Brief: D365 Copilot")
    assert "upload:brief.docx" in request["source_refs"]
    # extraction filled the fields from that text (provenance recorded)
    assert request["business_unit"] == "Technology"


def test_xlsx_upload_parses_sheets(client: TestClient) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Brief"
    sheet.append(["Objective", "Pipeline for D365 Copilot collateral offering"])
    sheet.append(["Vertical", "Manufacturing"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = _post_upload(
        client, "plan.xlsx", buffer,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert response.status_code == 200, response.text
    outcome = response.json()
    assert outcome["upload"]["kind"] == "xlsx"
    detail = client.get(f"/api/cases/{outcome['case_id']}").json()
    text = detail["summary"]["request"]["free_text_context"]
    assert "Objective | Pipeline for D365 Copilot" in text


def test_unsupported_upload_is_422_with_humane_message(client: TestClient) -> None:
    response = _post_upload(
        client, "deck.pptx", io.BytesIO(b"PK someslides"),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert response.status_code == 422
    assert ".docx" in response.json()["detail"]


def test_markdown_upload_runs_intake(client: TestClient) -> None:
    md = (
        "# Campaign Brief: D365 Copilot for Marketing Collateral\n\n"
        "Generate qualified pipeline among manufacturing accounts of $500M+ revenue."
    )
    response = _post_upload(client, "brief.md", io.BytesIO(md.encode("utf-8")), "text/markdown")
    assert response.status_code == 200, response.text
    outcome = response.json()
    assert outcome["upload"]["kind"] == "md"
    detail = client.get(f"/api/cases/{outcome['case_id']}").json()
    text = detail["summary"]["request"]["free_text_context"]
    assert text.startswith("# Campaign Brief: D365 Copilot")
