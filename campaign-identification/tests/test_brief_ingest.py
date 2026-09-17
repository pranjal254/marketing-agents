"""Uploaded-brief ingestion: faithful linearization, hard budgets, humane errors.
The parser never extracts fields itself — that stays with the L1 extraction layer."""

from __future__ import annotations

import io

import pytest

from campaign_identification.ingest import (
    MAX_UPLOAD_BYTES,
    TEXT_BUDGET_CHARS,
    BriefUploadError,
    parse_brief_document,
)


def _docx_bytes(*, paragraphs: int = 0) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Campaign Brief: D365 Copilot", level=1)
    document.add_paragraph("Objective: qualified pipeline among D365 customers.")
    document.add_heading("Timeline", level=2)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Phase"
    table.rows[0].cells[1].text = "Weeks 1-2"
    table.rows[1].cells[0].text = "Launch"
    table.rows[1].cells[1].text = "Landing page live"
    for i in range(paragraphs):
        document.add_paragraph(f"Filler paragraph {i} " + "x" * 120)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(*, rows: int = 3) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Brief"
    sheet.append(["Field", "Value"])
    sheet.append(["Business Unit", "Technology"])
    sheet.append(["Vertical", "Manufacturing"])
    for i in range(rows - 3):
        sheet.append([f"Row {i}", "y" * 80])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_docx_preserves_order_headings_and_tables() -> None:
    parsed = parse_brief_document("brief.docx", _docx_bytes())
    assert parsed.kind == "docx"
    lines = parsed.text.splitlines()
    assert lines[0] == "# Campaign Brief: D365 Copilot"
    assert "Objective: qualified pipeline among D365 customers." in lines[1]
    assert "## Timeline" in parsed.text
    assert "Launch | Landing page live" in parsed.text  # table rows as pipe lines
    assert parsed.tables == 1 and not parsed.truncated


def test_xlsx_rows_become_pipe_lines_per_sheet() -> None:
    parsed = parse_brief_document("plan.xlsx", _xlsx_bytes())
    assert parsed.kind == "xlsx"
    assert parsed.text.startswith("## Sheet: Brief")
    assert "Business Unit | Technology" in parsed.text
    assert parsed.paragraphs == 3 and parsed.tables == 1


def test_budget_truncates_never_fails() -> None:
    parsed = parse_brief_document("big.docx", _docx_bytes(paragraphs=400))
    assert parsed.truncated
    assert len(parsed.text) <= TEXT_BUDGET_CHARS
    assert parsed.text.startswith("# Campaign Brief")  # the head of the doc survives


def test_unsupported_type_and_empty_are_humane_errors() -> None:
    with pytest.raises(BriefUploadError, match=r"Word .*Markdown"):
        parse_brief_document("deck.pptx", b"PK someslides")
    with pytest.raises(BriefUploadError, match="empty"):
        parse_brief_document("brief.docx", b"")
    with pytest.raises(BriefUploadError, match="over 5 MB"):
        parse_brief_document("brief.docx", b"x" * (MAX_UPLOAD_BYTES + 1))


def test_corrupt_docx_is_a_humane_error() -> None:
    with pytest.raises(BriefUploadError, match="could not be read as a Word"):
        parse_brief_document("brief.docx", b"not a zip archive at all")


def _pdf_bytes(text: str) -> bytes:
    """A minimal valid one-page PDF with real text — no PDF-writing dependency."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_markdown_passes_through_with_budget() -> None:
    md = (
        "# Campaign Brief: D365 Copilot\n\n"
        "**Objective:** qualified pipeline.\n\n| Phase | Weeks 1-2 |"
    )
    parsed = parse_brief_document("brief.md", md.encode("utf-8"))
    assert parsed.kind == "md"
    assert parsed.text.splitlines()[0] == "# Campaign Brief: D365 Copilot"
    assert "| Phase | Weeks 1-2 |" in parsed.text
    assert not parsed.truncated


def test_pdf_text_is_extracted() -> None:
    parsed = parse_brief_document(
        "brief.pdf", _pdf_bytes("Campaign Brief: D365 Copilot pipeline for manufacturing")
    )
    assert parsed.kind == "pdf"
    assert "D365 Copilot pipeline for manufacturing" in parsed.text
    assert parsed.tables == 1  # pages read


def test_scanned_pdf_is_a_humane_error() -> None:
    with pytest.raises(BriefUploadError, match="looks scanned"):
        parse_brief_document("scan.pdf", _pdf_bytes(""))


def test_corrupt_pdf_is_a_humane_error() -> None:
    with pytest.raises(BriefUploadError, match="could not be read as a PDF"):
        parse_brief_document("brief.pdf", b"%PDF-1.4 then garbage with no objects")
